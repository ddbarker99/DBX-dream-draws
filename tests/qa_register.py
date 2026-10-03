"""Build the QA register (docs/qa/QA-REGISTER.csv and .xlsx) from real test runs.

    python -m pytest -q tests/test_app.py tests/test_qa_audit.py tests/test_security_audit.py --junitxml=/tmp/qa.xml
    python tests/qa_register.py /tmp/qa.xml [--extra docs/qa/extra_checks.json]

Columns: Test ID | Area | Feature | Test | Expected Result | Actual Result | Device/Browser | Pass/Fail | Severity |
Evidence | Bug ID | Retest Result.  Automated tests come from the JUnit XML (pass/fail is the real outcome); scripted
and manual checks (crawl, browser sweep, race, load, restore, items that need people or devices) come from the
--extra JSON, which records what was run and what wasn't.
"""
import ast
import csv
import json
import os
import re
import sys
import xml.etree.ElementTree as ET

ROOT = os.path.join(os.path.dirname(__file__), "..")
COLS = ["Test ID", "Area", "Feature", "Test", "Expected Result", "Actual Result", "Device/Browser", "Pass/Fail", "Severity",
        "Evidence", "Bug ID", "Retest Result"]

# Class -> (area, severity if this failed). P0 = money, entries, draws, security; P1 = core function; P2 = UX/validation.
AREAS = [
    (r"Money|Stripe|Payment|Wallet|Withdraw|Deposit|Refund|Finance|Checkout", "Payments & wallet", "P0"),
    (r"Draw|Postal|Snapshot|Instant|Game|Selection", "Competitions, draws & instant wins", "P0"),
    (r"Security|Admin|Role|Privilege|Csrf|Idor|Hardening|Auth|Session|Mfa|Isolation", "Security & access", "P0"),
    (r"Basket|Ticket|Allocation|Stampede|Concurren|Race|Oversell", "Ticket allocation & basket", "P0"),
    (r"Safer|Limit|Break|Exclusion|Responsible", "Responsible play", "P1"),
    (r"Recovery|Backup|Restore|Integrity|Ledger|Operations|Monitoring|Job|Reconcil", "Operations, integrity & recovery", "P1"),
    (r"Growth|Promo|Referral|Points|Experiment", "Promotions, points & referrals", "P1"),
    (r"Account|Customer|Phase6|Phase7|Trust|Overhaul|Evidence|Reporting|CreateForm", "Customer & admin experience", "P2"),
]
BUGS = {}


def area_for(cls, name):
    key = cls + " " + name
    for pat, area, sev in AREAS:
        if re.search(pat, key, re.I):
            return area, sev
    return "General", "P2"


def docstrings():
    """{(file, class, test): docstring} for every test function in tests/."""
    out = {}
    for fn in os.listdir(os.path.join(ROOT, "tests")):
        if not (fn.startswith("test_") and fn.endswith(".py")):
            continue
        tree = ast.parse(open(os.path.join(ROOT, "tests", fn), encoding="utf-8").read())
        for node in tree.body:
            if isinstance(node, ast.ClassDef):
                for f in node.body:
                    if isinstance(f, ast.FunctionDef) and f.name.startswith("test_"):
                        out[(fn[:-3], node.name, f.name)] = (ast.get_docstring(f) or "").strip()
    return out


def human(name):
    return name.replace("test_", "", 1).replace("_", " ").capitalize()


def main():
    xmls = [a for a in sys.argv[1:] if a.endswith(".xml")]
    extra = sys.argv[sys.argv.index("--extra") + 1] if "--extra" in sys.argv else os.path.join(ROOT, "docs", "qa", "extra_checks.json")
    bugs_path = os.path.join(ROOT, "docs", "qa", "bugs.json")
    bugs = json.load(open(bugs_path)) if os.path.exists(bugs_path) else []
    by_test = {}
    for b in bugs:
        for t in b.get("tests", []):
            by_test[t] = b["id"]
    docs = docstrings()
    rows, n = [], {}
    for x in xmls:
        for tc in ET.parse(x).getroot().iter("testcase"):
            module = tc.get("classname", "").rsplit(".", 1)
            mod, cls = (module[0].split(".")[-1], module[1]) if len(module) == 2 else ("", module[0])
            name = tc.get("name")
            failed = tc.find("failure") is not None or tc.find("error") is not None
            skipped = tc.find("skipped") is not None
            doc = docs.get((mod, cls, name), "")
            parts = [p.strip() for p in doc.split("|")] if doc.count("|") >= 3 else None
            area, sev = area_for(cls, name)
            if parts:
                tid, area, feature, expected = parts[0], parts[1], parts[2], parts[3]
            else:
                prefix = {"test_app": "REG", "test_security_audit": "SEC", "test_qa_audit": "QA"}.get(mod, "T")
                n[prefix] = n.get(prefix, 0) + 1
                tid, feature = f"{prefix}-{n[prefix]:03d}", cls
                expected = doc.split("\n")[0] if doc else human(name) + " — behaves as specified (assertions in the test)"
            msg = ""
            if failed:
                el = tc.find("failure") if tc.find("failure") is not None else tc.find("error")
                msg = (el.get("message") or "")[:200]
            bug = by_test.get(name, "")
            rows.append({"Test ID": tid, "Area": area, "Feature": feature, "Test": human(name), "Expected Result": expected,
                         "Actual Result": ("FAILED: " + msg) if failed else ("Skipped" if skipped else "As expected"),
                         "Device/Browser": "Automated (Flask test client)", "Pass/Fail": "Fail" if failed else ("Skip" if skipped else "Pass"),
                         "Severity": sev, "Evidence": f"tests/{mod}.py::{cls}::{name} ({tc.get('time')} s)", "Bug ID": bug,
                         "Retest Result": ("Pass after fix" if bug and not failed else ("" if not failed else "Open"))})
    if os.path.exists(extra):
        for r in json.load(open(extra)):
            rows.append({c: r.get(c, "") for c in COLS})
    os.makedirs(os.path.join(ROOT, "docs", "qa"), exist_ok=True)
    with open(os.path.join(ROOT, "docs", "qa", "QA-REGISTER.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "QA register"
        ws.append(COLS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7A1622")
        fills = {"Pass": "D9F2E3", "Fail": "F8D7DA", "Not run": "FFF3CD", "Skip": "EEEEEE"}
        for r in rows:
            ws.append([r[c] for c in COLS])
            colour = fills.get(r["Pass/Fail"])
            if colour:
                ws.cell(ws.max_row, COLS.index("Pass/Fail") + 1).fill = PatternFill("solid", fgColor=colour)
        for i, wdt in enumerate([12, 26, 26, 40, 60, 40, 24, 10, 9, 50, 9, 16], 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        bs = wb.create_sheet("Bugs")
        bcols = ["id", "severity", "area", "title", "found_by", "reproduce", "fix", "status", "retest", "tests"]
        bs.append(bcols)
        for b in bugs:
            bs.append([", ".join(b.get(k)) if isinstance(b.get(k), list) else b.get(k, "") for k in bcols])
        wb.save(os.path.join(ROOT, "docs", "qa", "QA-REGISTER.xlsx"))
    except ImportError:
        pass
    total = len(rows)
    by = lambda k: sum(1 for r in rows if r["Pass/Fail"] == k)  # noqa: E731
    print(f"rows={total} pass={by('Pass')} fail={by('Fail')} not_run={by('Not run')} skip={by('Skip')}")


if __name__ == "__main__":
    main()
